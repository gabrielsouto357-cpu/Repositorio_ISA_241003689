# -*- coding: utf-8 -*-
"""
Dicionário DEFINITIVO da base SINTÉTICA (Bloco 07): definição de cada campo + estatísticas observadas no arquivo de REFERÊNCIA gravado.
Saídas: DICIONARIO_DE_DADOS_SINTETICO.xlsx e DICIONARIO_DE_DADOS_SINTETICO.md (pasta base_B07).
Uso: python gerar_dicionario.py
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent; D = RAIZ / 'dados_sinteticos'
L = pd.read_csv(D / 'SINTETICO_lancamentos_2025_REFERENCIA.csv.gz', sep=';', dtype=str, keep_default_na=False)
S = pd.read_csv(D / 'SINTETICO_saldos_REFERENCIA.csv', sep=';', dtype=str, keep_default_na=False)
AV = 'SINTÉTICO DIDÁTICO – não são dados internos da ISA Energia Brasil. Totais por grupo reproduzem as DFs publicadas (F0003); a abertura em linhas é simulada.'
# campo: (tipo e unidade, descrição, real × sintético, regra de domínio / integridade)
F = {
 'id_lancamento': ('texto(11)', 'Chave primária do lançamento (L25 = arquivo 2025; L26 = arquivo jan–fev/2026), numerada em ordem cronológica', 'Sintético', 'Única; sem vazios; padrão L25/L26 + 8 dígitos'),
 'id_documento': ('texto', 'Documento-suporte. Prefixos: MED medição; AVC aviso de crédito/fatura ONS; REC aviso de recebimento; REM memória de remuneração; MEMO reclassificação; PROV provisão/estorno; PCLD perda esperada; MARCO marco contratual; RTP/EFI/REH atos ANEEL; SALDO abertura', 'Sintético (atos ANEEL citados são reais)', 'Repete só entre linhas do mesmo documento; (id_documento, seq_linha_documento) única'),
 'seq_linha_documento': ('inteiro', 'Sequência da linha no documento. Recebimento: linha 1 = parcela de O&M; linhas 2+ = ativo contratual', 'Sintético', '≥ 1; chave natural com id_documento'),
 'data_lancamento': ('data AAAA-MM-DD', 'Data do registro contábil', 'Sintético', '01/01–31/12/2025 no arquivo 2025; 01/01–28/02/2026 no arquivo 2026'),
 'hora_lancamento': ('hora HH:MM', 'Hora do registro; preenchida só em lançamentos manuais', 'Sintético', 'Vazia nos lotes automáticos e nos saldos; manual em horário comercial na referência'),
 'data_documento': ('data AAAA-MM-DD', 'Data do documento/evento econômico', 'Sintético', '≤ data_lancamento, exceto REMUNERACAO (cálculo no último dia do mês, lançado no último dia útil); 31/12/2024 só nos saldos de abertura'),
 'competencia': ('texto AAAA-MM', 'Mês a que o evento pertence', 'Sintético', 'Igual ao mês do lançamento, exceto RECEBIMENTO (mês da fatura paga) e saldos de abertura (2024-12)'),
 'entidade': ('texto', 'Entidade do grupo titular do contrato', 'Real (F0003 p.36)', 'Coerente com o contrato (cadastro REAL_contratos)'),
 'grupo': ('texto', 'Controladora ou Controlada', 'Real (F0003 p.36)', '{Controladora, Controlada}; coerente com o contrato'),
 'contrato': ('texto', 'Contrato de concessão ANEEL', 'Real (F0003 p.36)', 'FK para os 28 contratos consolidados'),
 'componente': ('texto', 'Subconta de Ativos da concessão (nota 5.3)', 'Real (F0003 p.45)', '{OEM, SE_FIN, SE_ECO, IMPL}'),
 'tipo_movimento': ('texto', 'Natureza do fluxo; ver aba Tipos_movimento', 'Linhas da nota 5.5 (F0003 p.47)', '13 valores do domínio (aba Tipos_movimento)'),
 'valor_movimento_brl': ('decimal R$, 2 casas, ponto decimal', 'FLUXO: + aumenta o ativo; − reduz. Σ de todas as linhas (com saldos de abertura) = saldo de 31/12/2025', 'Sintético; totais por grupo reais', 'Texto numérico exato (sem notação científica); sinal coerente com o tipo'),
 'id_obra_aneel': ('texto', 'Obra (IdeObr) da Lista de Módulos ANEEL 2025-2026', 'Real (F0242)', 'Vazio quando o módulo não tem obra no catálogo real (3.213 de 4.471 módulos)'),
 'id_modulo_aneel': ('texto', 'Módulo (IdeMdl) ANEEL da medição', 'Real (F0242)', 'Obrigatório em RECEITA_INFRA; FK ao catálogo; módulo do mesmo contrato'),
 'tipo_receita_aneel': ('texto', 'Tipo de receita do módulo', 'Real (F0242)', 'Domínio do catálogo (RBL, RBNI, RBNIA, RBSE, RCDM, RCDMA, RMEL, RPC, RPEC)'),
 'fornecedor_id': ('texto', 'Fornecedor do documento de custo', 'Sintético', 'F00001–F02500 (cadastro SINTETICO_cadastro_fornecedores); obrigatório em RECEITA_INFRA'),
 'tipo_custo': ('texto', 'Natureza do gasto de obra', 'Sintético', '7 categorias de obra; obrigatório em RECEITA_INFRA'),
 'custo_documento_brl': ('decimal R$', 'Custo do documento, sem margem', 'Sintético', 'valor_movimento = custo × (1 + margem), arredondado ao centavo'),
 'margem_construcao': ('decimal (fração)', 'Margem de construção do contrato', 'Sintético (arbitrado; CPC 47)', 'Faixa observada 0,020–0,077 (premissa 2%–9%)'),
 'usuario_pagador_id': ('texto', 'Usuário do sistema de transmissão (faturamento ONS)', 'Sintético', 'AG0001–AG0480; obrigatório em RECEITA_OEM e RECEBIMENTO'),
 'tipo_usuario': ('texto', 'Categoria do usuário', 'Sintético (categorias reais)', '{Distribuidora, Geradora, Consumidor livre, Exportador/Importador}'),
 'data_vencimento': ('data', 'Vencimento da fatura', 'Sintético', 'Dia 15 do mês seguinte; só RECEITA_OEM e RECEBIMENTO'),
 'dias_atraso': ('inteiro', 'Atraso do recebimento em dias', 'Sintético', '≥ 0; só RECEBIMENTO'),
 'taxa_implicita_aa': ('decimal (fração a.a.)', 'Taxa implícita do contrato usada na remuneração', 'Sintético dentro da faixa real 4,2%–9,9% (F0003 p.44)', 'Só REMUNERACAO do componente IMPL'),
 'base_calculo_brl': ('decimal R$', 'ESTOQUE de referência (saldo médio) usado no cálculo da remuneração do mês – NÃO SOMAR', 'Sintético', 'Informativo; nunca entra em totais nem em conciliação'),
 'origem': ('texto', 'AUTOMATICO (interfaces), MANUAL (lançamento humano) ou SALDO (abertura)', 'Sintético', '{AUTOMATICO, MANUAL, SALDO}'),
 'usuario_sistema': ('texto', 'Quem registrou: INT_* interface; SISTEMA; U01–U15 pessoas', 'Sintético', 'Manual: usuário pessoal'),
 'aprovador': ('texto', 'Aprovador: WF_* fluxo; SISTEMA; Uxx pessoa', 'Sintético', 'Manual: aprovador ≠ usuario_sistema (segregação)'),
 'historico': ('texto', 'Descrição do lançamento', 'Sintético', 'Texto livre padronizado por tipo'),
 'natureza_dado': ('texto', 'Marca obrigatória de dado sintético', 'Marca (F0002 §3)', 'Valor único SINTETICO_DIDATICO'),
}
FS = {'data_base': ('data', '31/12/2024 ou 31/12/2025', 'Somar só dentro da mesma data_base'), 'grupo': ('texto', 'Controladora ou Controlada', 'Coerente com o contrato'),
      'contrato': ('texto', 'Contrato (real, F0003 p.36)', 'FK'), 'componente': ('texto', 'Subconta', '{OEM, SE_FIN, SE_ECO, IMPL}'),
      'saldo_brl': ('decimal R$', 'ESTOQUE na data-base = abertura + Σ fluxos do contrato/componente', 'Σ por data_base = balanço (F0003 p.23)'),
      'parcela_circulante_brl': ('decimal R$', 'Parte realizável em até 12 meses (premissa: recebimentos dos 12 meses seguintes, calibrada)', 'Σ 2025 = 4.106.084 mil; O&M 100% circulante'),
      'parcela_nao_circulante_brl': ('decimal R$', 'Parte realizável após 12 meses', 'circulante + não circulante = saldo; ≥ 0'),
      'natureza_dado': ('texto', 'Marca obrigatória', 'SINTETICO_DIDATICO')}
TM = [('SALDO_ABERTURA', 'Saldo bruto em 31/12/2024 por contrato e componente', '+', 'Saldo de abertura'), ('SALDO_ABERTURA_PERDA_ESPERADA', 'Perda esperada acumulada em 31/12/2024 (redutora do O&M)', '−', 'Saldo de abertura'),
      ('RECEITA_INFRA', 'Item de medição de obra (custo + margem)', '+', 'Receita de infraestrutura'), ('PROVISAO_MEDICAO', 'Medição executada pendente de aprovação no fim do mês', '+', 'Receita de infraestrutura'),
      ('ESTORNO_PROVISAO', 'Estorno da provisão do mês anterior', '−', 'Receita de infraestrutura'), ('AJUSTE_MANUAL', 'Ajuste manual de obra (domínio previsto; sem ocorrência na referência)', '±', 'Receita de infraestrutura'),
      ('RECLASSIFICACAO', 'Transferência entre módulos do mesmo contrato; cada documento soma zero', '±', '—'), ('EFICIENCIA', 'Ganho/perda de eficiência na implementação', '±', 'Ganho (perda) de eficiência'),
      ('REMUNERACAO', 'Remuneração mensal do ativo contratual; inclui o ajuste RBSE da REH 3.466/2025', '±', 'Remuneração dos ativos de concessão'), ('RTP', 'Revisão tarifária periódica', '±', 'RTP'),
      ('RECEITA_OEM', 'Fatura mensal de O&M por usuário e contrato', '+', 'Receita de O&M'), ('PERDA_ESPERADA', 'Constituição de perda esperada (CPC 48)', '−', 'Perda esperada'),
      ('RECEBIMENTO', 'Pagamento de fatura; linha 1 O&M, linhas 2+ ativo contratual', '−', 'Recebimentos')]

def est(df, c):
    v = df[c]; nz = v[v != '']; dom = sorted(nz.unique())
    obs = ', '.join(dom[:8]) + (f' … ({len(dom)} valores)' if len(dom) > 8 else '') if len(dom) <= 40 or not nz.str.fullmatch(r'-?[\d.]+').all() else f'mín {min(nz, key=float)}; máx {max(nz, key=float)}'
    if c.startswith('data') or c == 'competencia' or c == 'hora_lancamento': obs = f'mín {nz.min()}; máx {nz.max()}' if len(nz) else '—'
    if c in ('id_lancamento', 'id_documento', 'historico'): obs = f'{len(dom):,} distintos'.replace(',', '.')
    return f'{100 * len(nz) / len(v):.1f}%', obs[:160]

wb = Workbook(); w0 = wb.active; w0.title = 'LEIA'
for i, t in enumerate([ 'DICIONÁRIO DE DADOS – BASE SINTÉTICA DE ATIVOS DA CONCESSÃO (versão definitiva – Bloco 07)', AV,
    'Formato: CSV UTF-8, separador ";" e PONTO decimal (valores gravados como texto exato a partir de centavos inteiros). Arquivos de lançamentos comprimidos em gzip.',
    'Mudança em relação à especificação do B06 (§4, que previa vírgula decimal): ponto decimal, para leitura sem conversão de localidade por pandas, R e Excel (importação com "." decimal). Não altera valores.',
    'Estatísticas observadas: arquivo SINTETICO_lancamentos_2025_REFERENCIA.csv.gz (600.650 linhas).',
    'Preparado por: matrícula 241003689 (execução assistida por IA). Revisão humana: ____________ (em branco).']):
    w0.cell(i + 1, 1, t)
w0.column_dimensions['A'].width = 170
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
def aba(nome, cab, linhas, larg):
    ws = wb.create_sheet(nome)
    for j, x in enumerate(cab, 1): c = ws.cell(1, j, x); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    for i, r in enumerate(linhas, 2):
        for j, x in enumerate(r, 1): ws.cell(i, j, x).alignment = Alignment(wrap_text=True, vertical='top')
    for j, w in enumerate(larg, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'
LF = [(k, *F[k][:2], F[k][2], F[k][3], *est(L, k)) for k in L.columns]
aba('Lancamentos_fluxos', ['Campo', 'Tipo e unidade', 'Descrição', 'Real × sintético', 'Domínio / regra de integridade', '% preenchido (REF)', 'Observado (REF)'], LF, [22, 18, 60, 26, 50, 12, 50])
LS = [(k, *FS[k], *est(S, k)) for k in S.columns]
aba('Saldos_estoque', ['Campo', 'Tipo', 'Descrição', 'Regra', '% preenchido', 'Observado'], LS, [26, 12, 60, 40, 12, 40])
cnt = L.tipo_movimento.value_counts()
aba('Tipos_movimento', ['tipo_movimento', 'Significado', 'Sinal', 'Linha da nota 5.5 (F0003 p.47)', 'Nº linhas (REF 2025)'], [(a, b, c, d, int(cnt.get(a, 0))) for a, b, c, d in TM], [30, 60, 8, 34, 16])
aba('Arquivos', ['Arquivo', 'Conteúdo', 'Natureza'], [
    ('dados_sinteticos/SINTETICO_lancamentos_2025_REFERENCIA.csv.gz', 'Lançamentos de 2025 + saldos de abertura – cenário sem erros, conciliado ao publicado', 'Sintético'),
    ('dados_sinteticos/SINTETICO_lancamentos_2025_TESTE.csv.gz', 'Mesma estrutura, com erros controlados (sem marcador): população dos testes', 'Sintético'),
    ('dados_sinteticos/SINTETICO_lancamentos_2026_jan_fev_*.csv.gz', 'Lançamentos posteriores (corte)', 'Sintético'),
    ('dados_sinteticos/SINTETICO_saldos_*.csv', 'Estoques por data-base, contrato e componente', 'Sintético (Σ = real)'),
    ('dados_sinteticos/SINTETICO_parametros_contratos.csv', 'Premissas arbitradas por contrato (taxa, margem, usuários, aberturas)', 'Sintético (RAP e contrato reais)'),
    ('dados_sinteticos/SINTETICO_cadastro_usuarios_pagadores.csv; SINTETICO_cadastro_fornecedores.csv', 'Dimensões', 'Sintético'),
    ('dados_reais_derivados/REAL_contratos_F0003_p36.csv', '28 contratos, entidade, grupo, RAP (REH 3.481/25), vencimento', 'Real – F0003 p.36'),
    ('dados_reais_derivados/REAL_catalogo_modulos_ANEEL_28_contratos.csv', '4.471 módulos ANEEL dos 28 contratos', 'Real – F0242'),
    ('gabarito_LACRADO/*', 'Registro de erros, linhas alteradas, mapa de IDs e ponte de efeito', 'LACRADO até o PT 450')], [70, 70, 30])
wb.save(RAIZ / 'DICIONARIO_DE_DADOS_SINTETICO.xlsx')
md = ['# Dicionário de dados – base sintética de Ativos da concessão (definitivo, Bloco 07)', '', f'> {AV}', '',
      '- Formato: CSV UTF-8, separador `;`, **ponto decimal**. Valores gravados como texto exato a partir de centavos inteiros. Lançamentos em gzip.',
      '- Mudança em relação ao B06 (§4 previa vírgula decimal): o ponto evita conversão de localidade. Os valores não mudam.', '',
      '## Lançamentos (fluxos) – 31 campos', '', '| Campo | Tipo | Descrição | Real × sintético | Regra | % preench. |', '|---|---|---|---|---|---|']
md += [f'| {a} | {b} | {c} | {d} | {e} | {f} |' for a, b, c, d, e, f, g in LF]
md += ['', '## Saldos (estoques) – 8 campos', '', '| Campo | Tipo | Descrição | Regra |', '|---|---|---|---|'] + [f'| {a} | {b} | {c} | {d} |' for a, b, c, d, e, f in LS]
md += ['', '## Tipos de movimento', '', '| Tipo | Significado | Sinal | Linha da nota 5.5 | Linhas REF |', '|---|---|---|---|---|'] + [f'| {a} | {b} | {c} | {d} | {int(cnt.get(a, 0)):,} |'.replace(',', '.') for a, b, c, d in TM]
open(RAIZ / 'DICIONARIO_DE_DADOS_SINTETICO.md', 'w', encoding='utf-8').write('\n'.join(md))
print('dicionário:', len(LF), 'campos de fluxo;', len(LS), 'de estoque;', len(TM), 'tipos')
