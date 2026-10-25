# -*- coding: utf-8 -*-
"""
Conciliação editável da base SINTÉTICA com as demonstrações publicadas (F0003 p.23, p.45, p.47).
Gera controle/CONCILIACAO_BASE_B07.xlsx (fórmulas vivas) e gabarito_LACRADO/PONTE_EFEITO_ERROS.xlsx.
Somas feitas em centavos inteiros (int64); a planilha recebe valores em R$ com 2 casas e converte para R$ mil por fórmula.
A conciliação do cenário de TESTE não usa o gabarito: mostra só a diferença agregada a investigar.
Uso: python conciliar.py
"""
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
CFG = json.load(open(RAIZ / 'config' / 'config_B07.json', encoding='utf-8')); M = CFG['metas_mil']
D = RAIZ / 'dados_sinteticos'; C = RAIZ / 'controle'; C.mkdir(exist_ok=True); G = RAIZ / 'gabarito_LACRADO'

def cent(s):
    s = s.astype(str); neg = s.str.startswith('-'); t = s.str.lstrip('-').str.split('.', expand=True)
    v = t[0].astype(np.int64) * 100 + t[1].str.ljust(2, '0').astype(np.int64)
    return pd.Series(np.where(neg, -v, v), index=s.index)
def ler(c):
    L = pd.read_csv(D / f'SINTETICO_lancamentos_2025_{c}.csv.gz', sep=';', dtype=str, keep_default_na=False)
    L['v'] = cent(L.valor_movimento_brl)
    S = pd.read_csv(D / f'SINTETICO_saldos_{c}.csv', sep=';', dtype=str, keep_default_na=False)
    for k in ['saldo_brl', 'parcela_circulante_brl', 'parcela_nao_circulante_brl']: S[k] = cent(S[k])
    return L, S
LINHA_55 = {'SALDO_ABERTURA': 'Saldo de abertura 31/12/2024', 'SALDO_ABERTURA_PERDA_ESPERADA': 'Saldo de abertura 31/12/2024',
            'RECEITA_INFRA': 'Receita de infraestrutura', 'PROVISAO_MEDICAO': 'Receita de infraestrutura', 'ESTORNO_PROVISAO': 'Receita de infraestrutura',
            'AJUSTE_MANUAL': 'Receita de infraestrutura', 'EFICIENCIA': 'Ganho (perda) de eficiência', 'REMUNERACAO': 'Remuneração dos ativos de concessão',
            'RECEITA_OEM': 'Receita de O&M', 'RTP': 'Revisão tarifária periódica (RTP)', 'PERDA_ESPERADA': 'Perda esperada',
            'RECEBIMENTO': 'Recebimentos', 'RECLASSIFICACAO': 'Reclassificações entre subcontas'}
ORDEM_55 = [('Saldo de abertura 31/12/2024', M['publicado_total']['2024']), ('Receita de infraestrutura', 6330048), ('Ganho (perda) de eficiência', -10513),
            ('Remuneração dos ativos de concessão', 2949237), ('Receita de O&M', 1264867), ('Revisão tarifária periódica (RTP)', 159203),
            ('Perda esperada', -4286), ('Recebimentos', -5304081), ('Reclassificações entre subcontas', 0)]
COMP = [('OEM', 'Serviços de O&M (ativo financeiro)', 158263), ('SE_FIN', 'Lei 12.783 – SE – componente financeiro', 3112255),
        ('SE_ECO', 'Lei 12.783 – SE – componente econômico', 3111847), ('IMPL', 'Implementação da infraestrutura', 30048933)]

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); B = Font(bold=True)
IN = PatternFill('solid', fgColor='FFF2CC'); RS = '#,##0.00;[Red]-#,##0.00'; RM = '#,##0;[Red]-#,##0'
def cab(ws, r, vals, larg=None):
    for j, x in enumerate(vals, 1):
        c = ws.cell(r, j, x); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    if larg:
        for j, w in enumerate(larg, 1): ws.column_dimensions[get_column_letter(j)].width = w
def bloco(ws, r0, titulo, linhas, nota='', total=''):
    """linhas: (rótulo, valor_base_centavos, publicado_mil, fonte). Colunas: A rótulo | B base R$ | C base R$ mil (fórmula) | D publicado R$ mil | E diferença | F fonte"""
    ws.cell(r0, 1, titulo).font = Font(bold=True, size=12); r = r0 + 1
    cab(ws, r, ['Item', 'Base sintética (R$)', 'Base (R$ mil) = B/1000', 'Publicado (R$ mil)', 'Diferença (R$ mil) = C − D', 'Fonte do publicado'])
    ini = r + 1
    for rot, vc, pub, fte in linhas:
        r += 1; ws.cell(r, 1, rot)
        ws.cell(r, 2, round(vc / 100, 2)).number_format = RS
        ws.cell(r, 3, f'=ROUND(B{r}/1000,3)').number_format = '#,##0.000;[Red]-#,##0.000'
        c = ws.cell(r, 4, pub); c.number_format = RM; c.fill = IN
        ws.cell(r, 5, f'=ROUND(C{r}-D{r},3)').number_format = '#,##0.000;[Red]-#,##0.000'; ws.cell(r, 6, fte)
    fim = r
    if total:
        r += 1; ws.cell(r, 1, total).font = B
        ws.cell(r, 2, f'=SUM(B{ini}:B{fim})').number_format = RS; ws.cell(r, 3, f'=ROUND(B{r}/1000,3)').number_format = '#,##0.000;[Red]-#,##0.000'
        ws.cell(r, 4, f'=SUM(D{ini}:D{fim})').number_format = RM; ws.cell(r, 5, f'=ROUND(C{r}-D{r},3)').number_format = '#,##0.000;[Red]-#,##0.000'
    if nota: r += 1; ws.cell(r, 1, nota).font = Font(italic=True)
    return r + 2, ini, fim

def conciliar(ws, L, S, cen):
    ws.cell(1, 1, f'Conciliação – cenário {cen} – base SINTÉTICA didática (não são dados internos da ISA Energia Brasil)').font = Font(bold=True, size=13)
    ws.cell(2, 1, 'Células amarelas = valores publicados digitados (F0003). Colunas C e E = fórmulas. Valores da base somados em centavos inteiros pelo script conciliar.py.')
    r = 4
    ctrl = L.grupo == 'Controladora'; s25 = S.data_base == '2025-12-31'; s24 = S.data_base == '2024-12-31'
    r, *_ = bloco(ws, r, 'C1 – Soma de todas as linhas (abertura + fluxos) = saldo de 31/12/2025',
                  [('Ativos da concessão – consolidado', int(L.v.sum()), M['publicado_total']['2025'], 'F0003 p.23 (4.106.084 + 32.325.214); p.45')])
    gc = L.groupby('componente').v.sum()
    r, a, b = bloco(ws, r, 'C2 – Saldo por componente (nota 5.3)', [(n, int(gc.get(k, 0)), p, 'F0003 p.45') for k, n, p in COMP], total='Total C2')
    r, *_ = bloco(ws, r, 'C3 – Subtotal ativo contratual (assunto principal de auditoria)',
                  [('SE financeiro + SE econômico + implementação', int(L[L.componente != 'OEM'].v.sum()), M['publicado_total']['ativo_contratual_2025'], 'F0003 p.45 e p.136 (KAM)')])
    g55 = L.assign(l=L.tipo_movimento.map(LINHA_55)).groupby('l').v.sum()
    r, a, b = bloco(ws, r, 'C4 – Movimentação do exercício por linha da nota 5.5 (consolidado)', [(n, int(g55.get(n, 0)), p, 'F0003 p.47') for n, p in ORDEM_55],
                    'Linhas por tipo_movimento: infraestrutura = RECEITA_INFRA + PROVISAO_MEDICAO + ESTORNO_PROVISAO (+ AJUSTE_MANUAL, se houver); remuneração inclui o ajuste RBSE da REH 3.466/2025 (−898.790).', total='Saldo final 31/12/2025 (Σ C4)')
    r, *_ = bloco(ws, r, 'C5 – Controladora (campo grupo)', [
        ('Saldo de abertura 31/12/2024 – controladora', int(L[ctrl & L.tipo_movimento.str.startswith('SALDO_ABERTURA')].v.sum()), M['publicado_total']['controladora_2024'], 'F0003 p.45'),
        ('Saldo 31/12/2025 – controladora', int(L[ctrl].v.sum()), M['publicado_total']['controladora_2025'], 'F0003 p.23, p.45'),
        ('Receita de infraestrutura – controladora', int(L[ctrl & L.tipo_movimento.map(LINHA_55).eq('Receita de infraestrutura')].v.sum()), M['IMPL']['Controladora']['infraestrutura'], 'F0003 p.47'),
        ('Receita de O&M – controladora', int(L[ctrl & (L.tipo_movimento == 'RECEITA_OEM')].v.sum()), M['OEM']['Controladora']['receita'], 'F0003 p.47')])
    r, *_ = bloco(ws, r, 'C6 – Tabela de saldos: circulante × não circulante', [
        ('Circulante 31/12/2025', int(S[s25].parcela_circulante_brl.sum()), M['circulante']['2025']['consolidado'], 'F0003 p.23'),
        ('Não circulante 31/12/2025', int(S[s25].parcela_nao_circulante_brl.sum()), M['publicado_total']['2025'] - M['circulante']['2025']['consolidado'], 'F0003 p.23'),
        ('Circulante 31/12/2024', int(S[s24].parcela_circulante_brl.sum()), M['circulante']['2024']['consolidado'], 'F0003 p.23'),
        ('Saldo 31/12/2024 (tabela de saldos)', int(S[s24].saldo_brl.sum()), M['publicado_total']['2024'], 'F0003 p.23')],
        'Saldos (estoques) só se somam dentro da mesma data-base; nunca somados aos fluxos.')
    # C7 – por contrato e componente
    ab = L[L.tipo_movimento.str.startswith('SALDO_ABERTURA')].groupby(['contrato', 'componente']).v.sum()
    fl = L[~L.tipo_movimento.str.startswith('SALDO_ABERTURA')].groupby(['contrato', 'componente']).v.sum()
    st = S[s25].groupby(['contrato', 'componente']).saldo_brl.sum(); s4 = S[s24].groupby(['contrato', 'componente']).saldo_brl.sum()
    ks = sorted(set(ab.index) | set(fl.index) | set(st.index))
    ws.cell(r, 1, 'C7 – Por contrato e componente: abertura (lançamentos) = saldo 2024 (tabela) e abertura + fluxos = saldo 2025 (tabela)').font = Font(bold=True, size=12); r += 1
    cab(ws, r, ['Contrato', 'Componente', 'Abertura – lançamentos (R$)', 'Saldo 31/12/2024 – tabela (R$)', 'Σ fluxos 2025 (R$)', 'Saldo 31/12/2025 – tabela (R$)', 'Dif. abertura = C − D', 'Dif. fechamento = C + E − F'])
    i0 = r + 1
    for k in ks:
        r += 1; vals = [k[0], k[1], ab.get(k, 0) / 100, s4.get(k, 0) / 100, fl.get(k, 0) / 100, st.get(k, 0) / 100]
        for j, x in enumerate(vals, 1):
            c = ws.cell(r, j, round(x, 2) if j > 2 else x)
            if j > 2: c.number_format = RS
        ws.cell(r, 7, f'=ROUND(C{r}-D{r},2)').number_format = RS; ws.cell(r, 8, f'=ROUND(C{r}+E{r}-F{r},2)').number_format = RS
    r += 1; ws.cell(r, 1, 'Totais').font = B
    for j in range(3, 9): ws.cell(r, j, f'=SUM({get_column_letter(j)}{i0}:{get_column_letter(j)}{r-1})').number_format = RS
    r += 1; ws.cell(r, 1, 'Pares contrato × componente com diferença ≠ 0'); ws.cell(r, 7, f'=COUNTIF(G{i0}:G{r-2},"<>0")'); ws.cell(r, 8, f'=COUNTIF(H{i0}:H{r-2},"<>0")')
    for j, w in enumerate([58, 20, 22, 20, 22, 30, 18, 18], 1): ws.column_dimensions[get_column_letter(j)].width = w
    return int(L.v.sum())

def subgrupos(ws, L, cen):
    ws.cell(1, 1, f'Valores por subgrupo – cenário {cen} (R$; base sintética). Tabela dinâmica estática: contrato × tipo de movimento').font = Font(bold=True, size=12)
    p = L.pivot_table(index=['grupo', 'contrato'], columns='tipo_movimento', values='v', aggfunc='sum', fill_value=0) / 100
    n = L.pivot_table(index=['grupo', 'contrato'], columns='tipo_movimento', values='v', aggfunc='count', fill_value=0)
    cols = list(p.columns); cab(ws, 3, ['Grupo', 'Contrato'] + cols + ['Total (fórmula)', 'Nº de linhas'])
    r = 3
    for (g, c), row in p.iterrows():
        r += 1; ws.cell(r, 1, g); ws.cell(r, 2, c)
        for j, x in enumerate(row.values, 3): ws.cell(r, j, round(float(x), 2)).number_format = RS
        ws.cell(r, 3 + len(cols), f'=SUM(C{r}:{get_column_letter(2 + len(cols))}{r})').number_format = RS
        ws.cell(r, 4 + len(cols), int(n.loc[(g, c)].sum()))
    r += 1; ws.cell(r, 1, 'Total').font = B
    for j in range(3, 5 + len(cols)): ws.cell(r, j, f'=SUM({get_column_letter(j)}4:{get_column_letter(j)}{r-1})').number_format = RS if j < 4 + len(cols) else '#,##0'
    ws.column_dimensions['A'].width = 13; ws.column_dimensions['B'].width = 11
    for j in range(3, 5 + len(cols)): ws.column_dimensions[get_column_letter(j)].width = 17

LR, SR = ler('REFERENCIA'); LT, ST = ler('TESTE')
wb = Workbook(); ws = wb.active; ws.title = 'LEIA'
txt = ['CONCILIAÇÃO DA BASE SINTÉTICA – Bloco 07 – matrícula 241003689 – execução assistida por IA – revisão humana: ____________ (em branco)',
       'Natureza: base SINTÉTICA para fins didáticos [SD]. Os totais por grupo reproduzem as demonstrações publicadas (F0003); a abertura em contratos, meses, linhas, usuários e fornecedores é simulada.',
       'Os registros NÃO são dados internos da ISA Energia Brasil e não demonstram nada sobre os registros reais da companhia.',
       'Abas: Ponte_unidades | C_REFERENCIA (C1–C7) | Subgrupos_REF | C_TESTE (mesmas conferências; diferença agregada a investigar) | Subgrupos_TESTE | Arredondamento.',
       'Precisão: somas em centavos inteiros (int64) no script; a planilha recebe R$ com 2 casas e converte para R$ mil com ROUND. Não há ajuste de fechamento.',
       'Cenário de TESTE: diferenças vêm de erros inseridos de propósito, documentados só no gabarito lacrado (gabarito_LACRADO/), que não é insumo desta conciliação nem dos testes; abrir só no PT 450.']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
ws.column_dimensions['A'].width = 160
# Ponte de unidades
wp = wb.create_sheet('Ponte_unidades'); totc = int(LR.v.sum())
cab(wp, 1, ['Passo', 'Valor', 'Unidade', 'Fórmula / origem'], [52, 24, 14, 70])
lin = [('Soma da base de referência', totc, 'centavos', 'Σ int64 de valor_movimento_brl × 100 (script)'),
       ('Em reais', '=B2/100', 'R$', 'B2/100'), ('Em milhares de reais', '=B3/1000', 'R$ mil', 'B3/1000 (escala das DFs)'),
       ('Publicado – balanço consolidado 31/12/2025', M['publicado_total']['2025'], 'R$ mil', 'F0003 p.23'),
       ('Diferença', '=ROUND(B4-B5,3)', 'R$ mil', 'Deve ser 0'), ('Moeda', 'BRL', '', 'Real; sem conversão cambial (operações em reais)')]
for i, (a, b, c, d) in enumerate(lin, 2):
    wp.cell(i, 1, a); x = wp.cell(i, 2, b); wp.cell(i, 3, c); wp.cell(i, 4, d)
    x.number_format = '#,##0.000' if i in (4, 5, 6) else '#,##0'
conciliar(wb.create_sheet('C_REFERENCIA'), LR, SR, 'REFERÊNCIA')
subgrupos(wb.create_sheet('Subgrupos_REF'), LR, 'REFERÊNCIA')
tt = conciliar(wb.create_sheet('C_TESTE'), LT, ST, 'TESTE')
w = wb['C_TESTE']; w.cell(3, 1, 'TESTE: a diferença em C1 é a "diferença a investigar" (razão sintético × balanço). Não é ajustada nem atribuída aqui; a explicação será obtida pelos testes (PT 500/520/530).').font = Font(bold=True, color='C00000')
subgrupos(wb.create_sheet('Subgrupos_TESTE'), LT, 'TESTE')
wa = wb.create_sheet('Arredondamento')
cab(wa, 1, ['Componente', 'Base (R$)', 'Base (R$ mil) sem arredondar', 'Base arredondada (R$ mil)', 'Publicado (R$ mil)', 'Resíduo de arredondamento (R$ mil)'], [44, 22, 24, 22, 18, 26])
gc = LR.groupby('componente').v.sum()
for i, (k, n, p) in enumerate(COMP, 2):
    wa.cell(i, 1, n); wa.cell(i, 2, gc[k] / 100).number_format = RS; wa.cell(i, 3, f'=B{i}/1000').number_format = '#,##0.000'
    wa.cell(i, 4, f'=ROUND(C{i},0)').number_format = RM; wa.cell(i, 5, p).number_format = RM; wa.cell(i, 6, f'=C{i}-E{i}').number_format = '#,##0.000'
wa.cell(7, 1, 'Total'); [wa.cell(7, j, f'=SUM({get_column_letter(j)}2:{get_column_letter(j)}5)') for j in range(2, 7)]
wa.cell(9, 1, 'Por desenho, cada total da base = publicado × 1.000 exatamente (alocação por maiores restos em centavos); o resíduo esperado é 0. Se as DFs têm arredondamento próprio entre linhas, ele não é recriado nem compensado na base.')
wb.save(C / 'CONCILIACAO_BASE_B07.xlsx')

# --------- ponte de efeito dos erros (LACRADA): referência -> cada erro -> teste
E = pd.read_csv(G / 'REGISTRO_ERROS.csv', sep=';', dtype=str, keep_default_na=False)
wb2 = Workbook(); w2 = wb2.active; w2.title = 'Ponte_efeito_erros'
w2.cell(1, 1, 'PONTE DE EFEITO DOS ERROS INSERIDOS – LACRADO: não usar como insumo de testes; abrir só no PT 450. Base sintética didática.').font = Font(bold=True, color='C00000')
cab(w2, 3, ['Item', 'Saldo 31/12/2025 (R$)', 'Resultado 2025 (R$)', 'Descrição', 'Nº linhas'], [40, 24, 24, 90, 10])
w2.cell(4, 1, 'Saldo da base de REFERÊNCIA (= publicado × 1.000)'); w2.cell(4, 2, totc / 100).number_format = RS
for i, e in enumerate(E.itertuples(), 5):
    w2.cell(i, 1, f'{e.id_erro} – {e.tipo}'); w2.cell(i, 2, float(e.efeito_saldo_311225_brl)).number_format = RS
    w2.cell(i, 3, float(e.efeito_resultado_2025_brl)).number_format = RS; w2.cell(i, 4, e.descricao); w2.cell(i, 5, int(e.n_linhas))
f = 5 + len(E); w2.cell(f, 1, 'Saldo esperado do TESTE (fórmula)').font = B; w2.cell(f, 2, f'=SUM(B4:B{f-1})').number_format = RS
w2.cell(f, 3, f'=SUM(C5:C{f-1})').number_format = RS
w2.cell(f + 1, 1, 'Saldo do TESTE somado do arquivo'); w2.cell(f + 1, 2, tt / 100).number_format = RS
w2.cell(f + 2, 1, 'Diferença não explicada (deve ser 0)').font = B; w2.cell(f + 2, 2, f'=ROUND(B{f}-B{f+1},2)').number_format = RS
w2.cell(f + 3, 1, 'Efeito líquido no saldo (R$ mil)'); w2.cell(f + 3, 2, f'=ROUND(SUM(B5:B{f-1})/1000,3)')
w2.cell(f + 4, 1, 'Efeito / ME (90.000)'); w2.cell(f + 4, 2, f'=B{f+3}/90000').number_format = '0.0%'
w2.cell(f + 6, 1, 'Sem compensação: nenhuma linha da referência foi reduzida para esconder efeito. E01, E08 e E10 não alteram o saldo e só se detectam por testes.')
wb2.save(G / 'PONTE_EFEITO_ERROS.xlsx')
print('REF (R$):', f'{totc/100:,.2f}', '| TESTE (R$):', f'{tt/100:,.2f}', '| diferença TESTE − publicado (R$ mil):', (tt - M['publicado_total']['2025'] * 100_000) / 100_000)
