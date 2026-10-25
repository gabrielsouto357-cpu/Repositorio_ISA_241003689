"""Reprodução acadêmica em cópia nova; nunca altera os arquivos entregues."""
import argparse, csv, gzip, hashlib, importlib.metadata, json, os
import platform, shutil, subprocess, sys
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT = Path(__file__).resolve().parent
BASE = Path('03_Base_sintetica/base_B07')
B8 = Path('04_Testes/B08')
VERSOES = {'numpy':'2.5.3','pandas':'3.0.5','scipy':'1.18.1','openpyxl':'3.1.5'}

def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1048576), b''): h.update(b)
    return h.hexdigest()

def comparar(a,b):
    op = gzip.open if a.suffix=='.gz' else open
    n=0; toleradas=0
    with op(a,'rt',encoding='utf-8',newline='') as f, op(b,'rt',encoding='utf-8',newline='') as g:
        aa=csv.reader(f,delimiter=';'); bb=csv.reader(g,delimiter=';')
        from itertools import zip_longest
        for x,y in zip_longest(aa,bb):
            n+=1
            if x is None or y is None or len(x)!=len(y): raise AssertionError(f'Granularidade divergente: {a.name}, linha {n}')
            for u,v in zip(x,y):
                if u==v: continue
                try:
                    du,dv=Decimal(u),Decimal(v)
                    # No máximo meia unidade de centavo, sem tolerância para identificadores.
                    if not ('.' in u or '.' in v) or not du.is_finite() or not dv.is_finite() or abs(du-dv)>Decimal('0.005'): raise InvalidOperation
                    toleradas+=1
                except InvalidOperation: raise AssertionError(f'Conteúdo divergente: {a.name}, linha {n}: {u!r} / {v!r}')
    return {'arquivo':a.relative_to(ROOT).as_posix(),'linhas_com_cabecalho':n,'diferencas_numericas_toleradas':toleradas,'aceite':True,'hash_identico':sha(a)==sha(b)}

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--saida',default='_reproducao/execucao_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    args=ap.parse_args()
    dest=(ROOT/args.saida).resolve()
    if dest.exists(): raise SystemExit('A pasta de saída já existe. Escolha outra com --saida. Nada foi sobrescrito.')
    if ROOT==dest or ROOT.is_relative_to(dest): raise SystemExit('Saída inválida: escolha uma subpasta nova.')
    env=os.environ.copy(); env['PYTHONUTF8']='1'; env['PYTHONDONTWRITEBYTECODE']='1'
    versões={k:importlib.metadata.version(k) for k in VERSOES}
    if versões!=VERSOES: raise SystemExit('Instale requirements.txt antes de continuar. Versões encontradas: '+str(versões))
    dest.mkdir(parents=True)
    report={'inicio_utc':datetime.now(timezone.utc).isoformat(),'python':sys.version,'sistema':platform.platform(),'versoes':versões,'etapas':[],'comparacoes':[],'aceite':False,
      'limite':'Reexecução automatizada por agente que conhece o gabarito. Não prova cegamento histórico, revisão independente ou evidência sobre a companhia.'}
    try:
        for rel in [BASE,B8,Path('04_Testes/B09'),Path('04_Testes/B12')]:
            shutil.copytree(ROOT/rel,dest/rel,ignore=shutil.ignore_patterns('__pycache__'))
        def run(rel):
            print('Executando '+str(rel),flush=True)
            r=subprocess.run([sys.executable,str(dest/rel)],cwd=dest,env=env,capture_output=True,text=True,encoding='utf-8',errors='replace')
            log=dest/('log_'+Path(rel).stem+'.txt'); log.write_text(r.stdout+'\n'+r.stderr,encoding='utf-8')
            report['etapas'].append({'programa':str(rel).replace('\\','/'),'retorno':r.returncode,'log':log.name})
            if r.returncode: raise RuntimeError('Falha em '+str(rel)+'; consultar '+log.name)
        run(BASE/'codigo/run_all.py')
        run(Path('04_Testes/B12/verificar_base_salva.py'))
        run(B8/'codigo/b08_testes.py')
        run(B8/'codigo/b08_td06_adicional.py')
        # O congelamento histórico original permanece intocado na entrega.
        cp=dest/B8/'controle/CONGELAMENTO_RESULTADOS.json'
        historical=json.loads(cp.read_text(encoding='utf-8'))
        shutil.copy2(cp,cp.with_name('CONGELAMENTO_HISTORICO_PRESERVADO.json'))
        checkpoint={'evento':'Checkpoint desta REEXECUÇÃO, antes do confronto; não substitui prova histórica de cegamento',
          'data_utc':datetime.now(timezone.utc).isoformat(),
          'arquivos':{f:sha(dest/B8/f) for f in historical['arquivos']}}
        cp.write_text(json.dumps(checkpoint,ensure_ascii=False,indent=2),encoding='utf-8')
        run(B8/'codigo/b08_confronto_gabarito.py')
        run(Path('04_Testes/B09/codigo/b09_politicas.py'))
        for sub in ['dados_sinteticos','dados_reais_derivados','gabarito_LACRADO']:
            for a in sorted((ROOT/BASE/sub).glob('*')):
                if a.name.endswith(('.csv','.csv.gz')): report['comparacoes'].append(comparar(a,dest/a.relative_to(ROOT)))
        for sub in ['resultados','amostras','confronto_PT450']:
            for a in sorted((ROOT/B8/sub).glob('*')):
                if a.name.endswith(('.csv','.csv.gz')): report['comparacoes'].append(comparar(a,dest/a.relative_to(ROOT)))
        a=ROOT/'04_Testes/B09/dados/COMPARACAO_POLITICAS_NOTA5.csv'
        report['comparacoes'].append(comparar(a,dest/a.relative_to(ROOT)))
        # Valores e fórmulas, sem comparar metadados ZIP nem caches de cálculo de XLSX.
        from openpyxl import load_workbook
        report['xlsx_comparados']=[]
        for rel in [BASE/'DICIONARIO_DE_DADOS_SINTETICO.xlsx',BASE/'controle/CONCILIACAO_BASE_B07.xlsx']:
            w1=load_workbook(ROOT/rel,data_only=False);w2=load_workbook(dest/rel,data_only=False)
            assert w1.sheetnames==w2.sheetnames, str(rel)
            for s in w1.sheetnames:
                assert list(w1[s].values)==list(w2[s].values), str(rel)+' / '+s
            report['xlsx_comparados'].append({'arquivo':rel.as_posix(),'valores_e_formulas_iguais':True})
        report['aceite']=True
        print('ACEITE: dados, amostra, detecções e confronto reproduzidos. Ver RELATORIO_REPRODUCAO.json.',flush=True)
    finally:
        report['fim_utc']=datetime.now(timezone.utc).isoformat()
        (dest/'RELATORIO_REPRODUCAO.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':main()
