# -*- coding: utf-8 -*-
"""PT 510 – compara o texto das políticas da nota 5 (5.1 e 5.2) entre DF 2025 (F0003) e DF 2024 (F0011).
Entrada: texto extraído (pdftotext) dos PDFs públicos; os trechos usados são gravados em dados/ para reprodução.
Saída: dados/politicas_nota5_F0003.txt, dados/politicas_nota5_F0011.txt, dados/COMPARACAO_POLITICAS_NOTA5.csv e aba Comparacao_texto no PT510."""
import re, csv, difflib, sys
from pathlib import Path
from openpyxl import load_workbook
B9 = Path(__file__).resolve().parent.parent
SRC = {'F0003': Path(sys.argv[1]) if len(sys.argv) > 1 else None, 'F0011': Path(sys.argv[2]) if len(sys.argv) > 2 else None}
def trecho(f, doc):
    if f is None or not f.exists():   # reprodução sem os PDFs: usa os trechos gravados
        return (B9 / f'dados/politicas_nota5_{doc}.txt').read_text(encoding='utf-8')
    pages = f.read_text(encoding='utf-8', errors='ignore').split('\f')
    t = '\n'.join(pages[40:47]); i = t.find('5.1.1'); j = t.find('5.3 ', i)
    j = t.find('Composição', i) if j < 0 else j
    return re.sub(r'\s+', ' ', t[i:j]).strip()
T = {k: trecho(v, k) for k, v in SRC.items()}
for k, v in T.items(): (B9 / f'dados/politicas_nota5_{k}.txt').write_text(v, encoding='utf-8')
sent = {k: [s.strip() for s in re.split(r'(?<=[\.;])\s+', v) if len(s.strip()) > 30] for k, v in T.items()}
ratio = difflib.SequenceMatcher(None, T['F0011'], T['F0003']).ratio()
so25 = [s for s in sent['F0003'] if max((difflib.SequenceMatcher(None, s, x).ratio() for x in sent['F0011']), default=0) < 0.85]
so24 = [s for s in sent['F0011'] if max((difflib.SequenceMatcher(None, s, x).ratio() for x in sent['F0003']), default=0) < 0.85]
with open(B9 / 'dados/COMPARACAO_POLITICAS_NOTA5.csv', 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh, delimiter=';'); w.writerow(['tipo', 'texto'])
    w.writerow(['similaridade_global', f'{ratio:.4f}'])
    for s in so25: w.writerow(['so_em_2025', s])
    for s in so24: w.writerow(['so_em_2024', s])
wb = load_workbook(B9 / 'planilhas/PT510_SALDOS_INICIAIS.xlsx')
if 'Comparacao_texto' in wb.sheetnames: del wb['Comparacao_texto']
ws = wb.create_sheet('Comparacao_texto'); ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 150
ws.append(['Item', 'Resultado (difflib; limiar de frase distinta: similaridade < 0,85)'])
ws.append(['Similaridade global', round(ratio, 4)]); ws.append(['Caracteres 2025 / 2024', f"{len(T['F0003'])} / {len(T['F0011'])}"])
for s in so25: ws.append(['Só em 2025', s])
for s in so24: ws.append(['Só em 2024', s])
wb.save(B9 / 'planilhas/PT510_SALDOS_INICIAIS.xlsx')
print(f'similaridade={ratio:.4f}; frases só 2025={len(so25)}; só 2024={len(so24)}')
for s in so25: print('2025>', s[:220])
for s in so24: print('2024>', s[:220])
