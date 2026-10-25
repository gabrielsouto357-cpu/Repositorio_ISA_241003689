# -*- coding: utf-8 -*-
"""Planilha editável do confronto (PT 450, parte simulada). Uso: python b08_confronto_planilha.py"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
B8 = Path(__file__).resolve().parent.parent
CF = pd.read_csv(B8 / 'confronto_PT450/CONFRONTO_GABARITO.csv', sep=';')
wb = Workbook(); ws = wb.active; ws.title = 'Confronto'
ws.cell(1, 1, 'Confronto testes × gabarito lacrado – base SINTÉTICA – desempenho DIDÁTICO; não é conclusão sobre a ISA Energia Brasil; mesma IA desenhou erros e testes (não é experimento cego)').font = Font(bold=True)
cab = ['Erro', 'Tipo', 'Linhas inseridas', 'Linhas sinalizadas', 'Linhas com distorção atribuída', 'Testes', 'Efeito saldo inserido (R$)', 'Efeito saldo detectado (R$)',
       'Diferença (R$)', 'Efeito resultado inserido (R$)', 'Efeito resultado detectado (R$)', 'Situação', 'Detectado? (1/0)']
for j, c in enumerate(cab, 1):
    x = ws.cell(3, j, c); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F3864'); x.alignment = Alignment(wrap_text=True)
for i, r in enumerate(CF.itertuples(), 4):
    for j, v in enumerate([r.id_erro, r.tipo, r.linhas_inseridas, r.linhas_sinalizadas, r.linhas_com_distorcao, r.testes, r.efeito_saldo_inserido, r.efeito_saldo_detectado], 1): ws.cell(i, j, v)
    ws.cell(i, 9, f'=ROUND(H{i}-G{i},2)'); ws.cell(i, 10, r.efeito_resultado_inserido); ws.cell(i, 11, r.efeito_resultado_detectado); ws.cell(i, 12, r.situacao)
    ws.cell(i, 13, f'=IF(LEFT(L{i},9)="Detectado",1,0)')
n = 3 + len(CF); t = n + 1
ws.cell(t, 1, 'Totais').font = Font(bold=True)
for c in 'CDEGHIJKM': ws[f'{c}{t}'] = f'=SUM({c}4:{c}{n})'
ws.cell(t + 2, 1, 'Tipos de erro detectados / inseridos'); ws.cell(t + 2, 3, f'=M{t}/COUNTA(A4:A{n})')
ws.cell(t + 3, 1, 'Linhas sinalizadas / inseridas'); ws.cell(t + 3, 3, f'=D{t}/C{t}')
ws.cell(t + 4, 1, 'Efeito no saldo detectado / inserido'); ws.cell(t + 4, 3, f'=H{t}/G{t}')
ws.cell(t + 5, 1, 'Falsos positivos com distorção atribuída (linhas)'); ws.cell(t + 5, 3, 0)
ws.cell(t + 6, 1, 'Exceções examinadas sem erro: 35 primeiras ocorrências de documentos repetidos (corretamente sem distorção) e 1 sinal TD-06 (059/2001 SE_FIN) explicado por evidência pública')
for c, w in zip('ABCDEFGHIJKLM', [7, 26, 10, 10, 12, 22, 18, 18, 12, 18, 18, 36, 10]): ws.column_dimensions[c].width = w
wb.save(B8 / 'confronto_PT450/CONFRONTO_PT450.xlsx'); print('ok')
