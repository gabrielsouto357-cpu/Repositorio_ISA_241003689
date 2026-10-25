# -*- coding: utf-8 -*-
"""Memória de cálculo editável do Bloco 08 (fórmulas vivas). Lê só os resultados do b08_testes.py (sem gabarito). Uso: python b08_memoria.py"""
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
B8 = Path(__file__).resolve().parent.parent; R = B8 / 'resultados'
X = json.load(open(R / 'RESULTADOS_B08.json', encoding='utf-8')); M = X['mus']; E = X['estagio1']
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); IN = PatternFill('solid', fgColor='FFF2CC'); B = Font(bold=True)
RS = '#,##0.00;[Red]-#,##0.00'
wb = Workbook(); ws = wb.active; ws.title = 'LEIA'
for i, t in enumerate(['MEMÓRIA DE CÁLCULO – BLOCO 08 (PT 530 e parte simulada do PT 450) – base SINTÉTICA, cenário TESTE – não são dados da ISA Energia Brasil',
                       'Células amarelas = entradas (parâmetros do desenho congelado ou resultados do b08_testes.py). Demais = fórmulas. Valores em R$.',
                       'Fator de confiança: GAMMAINV(0,95; k+1; 1) = limite superior de Poisson para k distorções (equivale à tabela usual de MUS).',
                       'Sem uso do gabarito: esta planilha foi gerada antes do confronto (controle/CONGELAMENTO_RESULTADOS.json).',
                       'Preparado por: matrícula 241003689 (execução assistida por IA). Revisão humana: ____________ (em branco).'], 1): ws.cell(i, 1, t)
ws.column_dimensions['A'].width = 150
def cab(w, r, vals, larg):
    for j, v in enumerate(vals, 1): c = w.cell(r, j, v); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    for j, l in enumerate(larg, 1): w.column_dimensions[get_column_letter(j)].width = l
def linha(w, r, rot, val, fmt=RS, entrada=False, nota=''):
    w.cell(r, 1, rot); c = w.cell(r, 2, val); c.number_format = fmt
    if entrada: c.fill = IN
    w.cell(r, 3, nota)
# ---------- Dimensionamento
d = wb.create_sheet('Dimensionamento_MUS'); cab(d, 1, ['Parâmetro', 'Valor', 'Fonte / natureza'], [62, 22, 80])
L = [('Materialidade de execução – ME (R$)', 90_000_000, RS, True, 'D-014 – exigência: TM ≤ ME (530 A3)'),
     ('Σ |efeito no saldo| – Estágio 1 (R$)', E['soma_abs_efeito_saldo_brl'], RS, True, 'Resultado b08_testes.py'),
     ('Σ |efeito no resultado| – Estágio 1 (R$)', E['soma_abs_efeito_resultado_brl'], RS, True, 'Resultado b08_testes.py'),
     ('TM da amostra (R$) = ME − máx', '=B2-MAX(B3,B4)', RS, False, 'Regra do desenho (julgamento conservador: distorções presumidas não corrigidas)'),
     ('Erro esperado – EE (R$)', 9_000_000, RS, True, 'Premissa didática (B06; 10% da ME) – julgamento'),
     ('Fator de expansão – FE', 1.6, '0.00', True, 'Tabela usual de MUS, 5% de risco – julgamento'),
     ('Confiança', 0.95, '0%', True, 'Risco significativo, controles não testados (PT 330) – julgamento'),
     ('Fator de confiança 0 erros – FC(0)', '=GAMMAINV(B8,1,1)', '0.0000', False, 'Poisson: −ln(0,05)'),
     ('Valor contábil do restante – VC (R$)', M['VC_restante_brl'], RS, True, 'População TD-01 − itens específicos − estrato superior'),
     ('Tamanho n = ⌈VC × FC / (TM − EE × FE)⌉', '=ROUNDUP(B10*B9/(B5-B6*B7),0)', '0', False, '530 it. 7; A11 (fórmula estatística)'),
     ('Intervalo J (R$) = ⌊VC / n⌋', '=ROUNDDOWN(B10*100/B11,0)/100', RS, False, 'Itens ≥ J vão para o estrato superior (100%)'),
     ('Início aleatório (centavos)', M['inicio_aleatorio_centavos'], '0', True, 'Semente 2025123108 (numpy PCG64)'),
     ('n efetivamente selecionado (código)', M['n_selecionado'], '0', True, 'AMOSTRA_MUS.csv'),
     ('Conferência: n fórmula − n código', '=B11-B14', '0', False, 'Deve ser 0'),
     ('Referência B06 (sem Estágio 1; VC total; TM = ME)', X['tamanho_referencia_B06'], '0', True, 'Comparação – tamanho maior agora porque o TM foi reduzido pelas distorções já encontradas')]
for i, (a, b, f, e, n) in enumerate(L, 2): linha(d, i, a, b, f, e, n)
# ---------- Avaliação
a = wb.create_sheet('Avaliacao_MUS'); cab(a, 1, ['Item', 'Valor', 'Fórmula / nota'], [62, 22, 80])
am = pd.read_csv(B8 / 'amostras' / 'AMOSTRA_MUS.csv', sep=';'); tt = sorted(am.tainting[am.tainting > 0], reverse=True)
L = [('J (R$)', "=Dimensionamento_MUS!B12", RS, False, ''), ('Confiança', '=Dimensionamento_MUS!B8', '0%', False, ''),
     ('Itens na amostra', len(am), '0', True, ''), ('Itens com distorção na amostra', len(tt), '0', True, ''),
     ('Soma das distorções detectadas na amostra (R$)', float(am.distorcao_centavos.sum()) / 100, RS, True, 'Valor real encontrado – NÃO é a projeção'),
     ('Precisão básica = FC(0) × J', '=GAMMAINV(B3,1,1)*B2', RS, False, 'Risco de amostragem com 0 distorções'),
     ('Distorção projetada = Σ t × J', f'=SUM(B{len(tt) and 20}:B{19 + max(len(tt), 1)})*B2' if tt else 0, RS, False, '530 it. 14; A18'),
     ('Acréscimo por distorções ordenadas', f'=SUMPRODUCT(B20:B{19 + len(tt)},C20:C{19 + len(tt)})*B2' if tt else 0, RS, False, 'Σ t(i) × (FC(i) − FC(i−1) − 1)'),
     ('LIMITE SUPERIOR de distorção (restante)', '=B7+B8+B9', RS, False, 'Limite de avaliação (inclui risco de amostragem)'),
     ('TM da amostra', '=Dimensionamento_MUS!B5', RS, False, ''),
     ('Folga = TM − limite superior', '=B11-B10', RS, False, '> 0: amostra dá base razoável (530 A22)'),
     ('Estrato superior: itens / valor / distorção real (R$)', f"{M['estrato_superior_itens']} / {M['estrato_superior_valor_brl']:,.2f} / {M['estrato_superior_distorcao_brl']:,.2f}", '@', True, 'Examinado a 100%; sem projeção')]
for i, (x, b, f, e, n) in enumerate(L, 2): linha(a, i, x, b, f, e, n)
a.cell(19, 1, 'Contaminações ordenadas (t) e incremento de fator').font = B
for k, t in enumerate(tt):
    a.cell(20 + k, 2, t); a.cell(20 + k, 3, f'=GAMMAINV($B$3,{k + 2},1)-GAMMAINV($B$3,{k + 1},1)-1')
if not tt: a.cell(20, 1, 'Nenhuma distorção na amostra: distorção projetada = 0; limite superior = precisão básica.')
# ---------- Estágio 1
s = wb.create_sheet('Estagio1_distorcoes'); cab(s, 1, ['Teste | regra', 'Linhas', 'Efeito no saldo 31/12/2025 (R$)', 'Efeito no resultado 2025 (R$)', 'Natureza (450 A6)', 'Meio de seleção'], [22, 10, 26, 26, 22, 40])
nat = {'TD-01|R1': 'Factual', 'TD-01|R2': 'Factual', 'TD-01|R4': 'Factual', 'TD-01|R5': 'Factual (corte; afeta abertura)', 'TD-01|R6': 'Factual (sem suporte; indício de fraude)',
       'TD-03|CORTE-2026': 'Factual (subavaliação)', 'TD-05|REM': 'Factual (cálculo)', 'TD-07|ESTORNO': 'Factual'}
i = 1
for k, v in E['por_teste'].items():
    i += 1; s.cell(i, 1, k); s.cell(i, 2, v['linhas']); s.cell(i, 3, v['saldo_brl']).number_format = RS; s.cell(i, 4, v['resultado_brl']).number_format = RS
    s.cell(i, 5, nat.get(k, 'Factual')); s.cell(i, 6, '100% / itens específicos (não é amostragem; sem projeção – 500 A66)')
i += 1; s.cell(i, 1, 'Total líquido').font = B
for c in (2, 3, 4): s.cell(i, c, f'=SUM({get_column_letter(c)}2:{get_column_letter(c)}{i - 1})').number_format = RS if c > 2 else '0'
s.cell(i + 1, 1, 'Total |absoluto|').font = B; s.cell(i + 1, 3, f'=SUMPRODUCT(ABS(C2:C{i - 1}))').number_format = RS; s.cell(i + 1, 4, f'=SUMPRODUCT(ABS(D2:D{i - 1}))').number_format = RS
s.cell(i + 3, 1, 'Classificação (sem efeito no total): TD-04 recebimentos aplicados a outro contrato (R$)'); s.cell(i + 3, 3, E['classificacao_contratos_TD04_brl']).number_format = RS
# ---------- Agregação PT 450
g = wb.create_sheet('Agregacao_PT450'); cab(g, 1, ['Componente', 'Saldo 31/12/2025 (R$)', 'Resultado 2025 (R$)', 'Nota'], [62, 24, 24, 80])
n1 = i
rows = [('(1) Distorções factuais identificadas – líquido', f'=Estagio1_distorcoes!C{n1}', f'=Estagio1_distorcoes!D{n1}', '100%/itens específicos; sem risco de amostragem'),
        ('(2) Distorções de julgamento', 0, 0, 'TD-06: exceção explicada por evidência pública (F0003 p.37, 46); nenhuma registrada'),
        ('(3) Distorção projetada (MUS, restante)', '=Avaliacao_MUS!B8', '=Avaliacao_MUS!B8', 'Superavaliação de débitos de receita/ativo'),
        ('(4) Melhor estimativa = (1)+(2)+(3)', '=B2+B3+B4', '=C2+C3+C4', '450 A6; 530 A22'),
        ('(5) Margem para risco de amostragem = limite superior − projetada', '=Avaliacao_MUS!B10-Avaliacao_MUS!B8', '=Avaliacao_MUS!B10-Avaliacao_MUS!B8', 'Não é distorção identificada'),
        ('(6) Limite superior agregado = (4)+(5)', '=B5+B6', '=C5+C6', 'Comparar com ME e MG'),
        ('ME', 90_000_000, 90_000_000, 'D-014'), ('MG', 139_000_000, 139_000_000, 'D-014'), ('CT', 7_000_000, 7_000_000, 'D-014'),
        ('(6) / ME', '=B7/B8', '=C7/C8', ''),
        ('Conciliação: TESTE − publicado (PT 500, C1)', 31_725_784.44, None, 'CONCILIACAO_BASE_B07 aba C_TESTE'),
        ('Explicado pelas distorções factuais líquidas (saldo)', '=B2', None, ''),
        ('Diferença de conciliação NÃO explicada', '=ROUND(B12-B13,2)', None, 'Localizada em O&M / perda esperada (C2 e C4 do C_TESTE): a investigar no PT 540')]
for k, (x, b, c, nn) in enumerate(rows, 2):
    g.cell(k, 1, x); g.cell(k, 2, b).number_format = RS if k != 11 else '0.0%'
    if c is not None: g.cell(k, 3, c).number_format = RS if k != 11 else '0.0%'
    g.cell(k, 4, nn)
for k in (8, 9, 10): g.cell(k, 2).fill = IN; g.cell(k, 3).fill = IN
g.cell(12, 2).fill = IN
# ---------- TD-05 e TD-06
t5 = pd.read_csv(R / 'TD05_REMUNERACAO_RECALCULO.csv', sep=';'); w5 = wb.create_sheet('TD05_remuneracao')
cab(w5, 1, ['Contrato', 'Comp.', 'Competência', 'Base (R$)', 'Valor (R$)', 'Taxa', 'Taxa modal', 'Fator modal', 'Esperado = base × fator', 'Diferença'], [10, 8, 11, 18, 16, 8, 10, 12, 20, 14])
for k, r in enumerate(t5.itertuples(), 2):
    for j, v in enumerate([r.contrato, r.componente, r.competencia, r.base_brl, r.valor_brl, r.taxa, r.taxa_modal, r.fator_modal], 1): w5.cell(k, j, v)
    w5.cell(k, 9, f'=ROUND(D{k}*H{k},2)').number_format = RS; w5.cell(k, 10, f'=ROUND(E{k}-I{k},2)').number_format = RS
nk = len(t5) + 2; w5.cell(nk, 1, 'Σ |diferenças| > R$ 1'); w5.cell(nk, 10, f'=SUMPRODUCT((ABS(J2:J{nk - 1})>1)*J2:J{nk - 1})').number_format = RS
t6 = pd.read_csv(R / 'TD06_CLASSIFICACAO.csv', sep=';'); w6 = wb.create_sheet('TD06_classificacao')
cab(w6, 1, ['Contrato', 'Comp.', 'Índice 2024', 'Índice 2025', 'Circulante 2025', 'Esperado (índice 2024)', 'Diferença', 'Variação p.p.', 'Exceção pela regra'], [10, 8, 11, 11, 18, 20, 18, 12, 14])
for k, r in enumerate(t6.itertuples(), 2):
    for j, v in enumerate([r.contrato, r.componente, r.indice_2024, r.indice_2025, r.circulante_2025_brl, r.circulante_esperado_brl], 1): w6.cell(k, j, v)
    w6.cell(k, 7, f'=E{k}-F{k}').number_format = RS; w6.cell(k, 8, f'=(D{k}-C{k})*100').number_format = '0.0'
    w6.cell(k, 9, f'=IF(OR(AND(B{k}="OEM",ABS(D{k}-1)>0.000001),AND(ABS(H{k})>10,ABS(G{k})>=7000000)),"SIM","não")')
wb.save(R / 'MEMORIA_CALCULO_B08.xlsx'); print('ok')
